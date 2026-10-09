"""Generate an idempotent SQL import from the approved project portfolio workbook.

Only fields rendered by the portfolio or project-detail pages are imported. Image
URLs are deliberately excluded because the workbook does not supply approved card
images; an existing CMS image is therefore never overwritten.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from openpyxl import load_workbook


FIELDS = (
    "title, slug, match_slugs, summary, description, project_status, project_year, "
    "impact_area, work_stream, capabilities_involved, best_fit, problem, solution, "
    "how_it_works, features, tech_stack, collaboration_network, "
    "implementation_countries, resource_links, video_url, display_order"
)

COLUMN_DEFINITIONS = """
  title text NOT NULL,
  slug text NOT NULL,
  match_slugs text[] NOT NULL DEFAULT ARRAY[]::text[],
  summary text NOT NULL,
  description text NOT NULL,
  project_status text NOT NULL,
  project_year integer,
  impact_area text NOT NULL,
  work_stream text NOT NULL,
  capabilities_involved text[] NOT NULL DEFAULT ARRAY[]::text[],
  best_fit text[] NOT NULL DEFAULT ARRAY[]::text[],
  problem text NOT NULL,
  solution text NOT NULL,
  how_it_works text[] NOT NULL DEFAULT ARRAY[]::text[],
  features text[] NOT NULL DEFAULT ARRAY[]::text[],
  tech_stack text[] NOT NULL DEFAULT ARRAY[]::text[],
  collaboration_network text NOT NULL,
  implementation_countries text[] NOT NULL DEFAULT ARRAY[]::text[],
  resource_links text[] NOT NULL DEFAULT ARRAY[]::text[],
  video_url text,
  display_order integer NOT NULL
""".strip()


def clean(value: object) -> str:
    return str(value or "").replace("�", "-").strip()


def slugify(value: str) -> str:
    value = clean(value).lower()
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value


def sql_text(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def sql_array(values: list[str]) -> str:
    if not values:
        return "ARRAY[]::text[]"
    return "ARRAY[" + ", ".join(sql_text(value) for value in values) + "]::text[]"


def list_value(value: object, split_inline: bool = False) -> list[str]:
    text = clean(value)
    if not text:
        return []
    if split_inline and "\n" not in text and "\r" not in text:
        text = re.sub(r"\s*[;,]\s*", "\n", text)
    return [
        re.sub(r"^[•*\-]+\s*", "", line).strip()
        for line in re.split(r"[\r\n]+", text)
        if line.strip()
    ]


def links(value: object) -> list[str]:
    return re.findall(r"https?://[^\s,;]+", clean(value))


def impact_area(work_stream: str) -> str:
    value = work_stream.lower()
    if "gis" in value:
        return "GIS / Remote Sensing"
    if "nlp" in value:
        return "Natural Language Processing"
    if "training" in value:
        return "Digital Skills Development"
    if "resilience" in value or "open-source" in value:
        return "Resilience"
    return ""


def project_status(value: str) -> str:
    return "completed" if value.lower() == "completed" else "active"


def main() -> None:
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_project_portfolio_import.py <portfolio.xlsx>")

    workbook = load_workbook(Path(sys.argv[1]), read_only=True, data_only=True)
    sheet = workbook["Product Portfolio"]
    header = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
    rows = [dict(zip(header, row)) for row in sheet.iter_rows(min_row=7, values_only=True)]

    values: list[str] = []
    for order, row in enumerate((row for row in rows if clean(row.get("Project Name "))), start=1):
        title = clean(row["Project Name "])
        slug = slugify(title)
        # The legacy CMS used this established slug for the PFS project. Keeping
        # it as a match key updates that record rather than creating a duplicate.
        match_slugs = ["public-finance-simplification-platform-pfs"] if slug == "public-finance-simplification-app" else []
        general = clean(row.get("Project description (general)"))
        work_stream = clean(row.get("Work stream"))
        video_links = links(row.get("Short video / media"))
        year = clean(row.get("Year "))
        values.append(
            "(" + ", ".join((
                sql_text(title),
                sql_text(slug),
                sql_array(match_slugs),
                sql_text(general),
                sql_text(general),
                sql_text(project_status(clean(row.get("Status")))),
                year if year.isdigit() else "NULL",
                sql_text(impact_area(work_stream)),
                sql_text(work_stream),
                sql_array(list_value(row.get("Capabilities involved"), split_inline=True)),
                sql_array(list_value(row.get("Current client segments served"))),
                sql_text(clean(row.get("Problem"))),
                sql_text(clean(row.get("Solution"))),
                sql_array(list_value(row.get("How it works"))),
                sql_array(list_value(row.get("Features"))),
                sql_array(list_value(row.get("Tech stack"), split_inline=True)),
                sql_text(clean(row.get("Collaboration / our network"))),
                sql_array(list_value(row.get("Implementation country"), split_inline=True)),
                sql_array(links(row.get("Downloads"))),
                sql_text(video_links[0]) if video_links else "NULL",
                str(order),
            )) + ")"
        )

    print("-- Generated from the approved UNDP Project Portfolio workbook. Do not edit generated values by hand.\n")
    print("BEGIN;\n")
    print("CREATE TEMP TABLE portfolio_import (\n" + COLUMN_DEFINITIONS + "\n) ON COMMIT DROP;\n")
    print("INSERT INTO portfolio_import (" + FIELDS + ") VALUES")
    print(",\n".join(values) + ";\n")
    print("""
WITH matched AS (
  SELECT source.*, target.id
  FROM portfolio_import source
  JOIN LATERAL (
    SELECT id
    FROM projects
    WHERE slug = source.slug
      OR slug = ANY(source.match_slugs)
      OR lower(regexp_replace(title, '[^a-z0-9]+', '', 'g')) =
         lower(regexp_replace(source.title, '[^a-z0-9]+', '', 'g'))
    ORDER BY (slug = source.slug) DESC, updated_at DESC
    LIMIT 1
  ) target ON true
)
UPDATE projects target
SET
  title = source.title,
  summary = source.summary,
  description = source.description,
  project_status = source.project_status,
  project_year = source.project_year,
  impact_area = NULLIF(source.impact_area, ''),
  work_stream = NULLIF(source.work_stream, ''),
  capabilities_involved = source.capabilities_involved,
  best_fit = source.best_fit,
  problem = NULLIF(source.problem, ''),
  solution = NULLIF(source.solution, ''),
  how_it_works = source.how_it_works,
  features = source.features,
  tech_stack = source.tech_stack,
  collaboration_network = NULLIF(source.collaboration_network, ''),
  implementation_countries = source.implementation_countries,
  resource_links = source.resource_links,
  video_url = COALESCE(source.video_url, target.video_url),
  display_order = source.display_order,
  status = 'published',
  published_at = COALESCE(target.published_at, now()),
  updated_at = now()
FROM matched source
WHERE target.id = source.id;

INSERT INTO projects (
  title, slug, summary, description, project_status, project_year, impact_area,
  work_stream, capabilities_involved, best_fit, problem, solution, how_it_works,
  features, tech_stack, collaboration_network, implementation_countries,
  resource_links, video_url, display_order, status, published_at
)
SELECT
  source.title, source.slug, source.summary, source.description,
  source.project_status, source.project_year, NULLIF(source.impact_area, ''),
  NULLIF(source.work_stream, ''), source.capabilities_involved, source.best_fit,
  NULLIF(source.problem, ''), NULLIF(source.solution, ''), source.how_it_works,
  source.features, source.tech_stack, NULLIF(source.collaboration_network, ''),
  source.implementation_countries, source.resource_links, source.video_url,
  source.display_order, 'published', now()
FROM portfolio_import source
WHERE NOT EXISTS (
  SELECT 1
  FROM projects target
  WHERE target.slug = source.slug
     OR target.slug = ANY(source.match_slugs)
     OR lower(regexp_replace(target.title, '[^a-z0-9]+', '', 'g')) =
        lower(regexp_replace(source.title, '[^a-z0-9]+', '', 'g'))
);

SELECT
  COUNT(*) FILTER (WHERE status = 'published') AS published_projects,
  COUNT(*) FILTER (WHERE status = 'draft') AS draft_projects,
  COUNT(*) FILTER (WHERE status = 'archived') AS archived_projects
FROM projects;

COMMIT;
""")


if __name__ == "__main__":
    main()
